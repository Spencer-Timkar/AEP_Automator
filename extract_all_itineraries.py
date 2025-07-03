from PyQt5.QtWidgets import QApplication, QFileDialog
import sys
import os
import re
from openpyxl import load_workbook
import pandas as pd
import matplotlib.pyplot as plt
import datetime


# Function to select itinerary folder using a file dialog
def get_itinerary_folder():
    app = QApplication(sys.argv)
    folder = QFileDialog.getExistingDirectory(None, "Select Folder Containing Itineraries")
    app.exit()
    return folder

class itinerary:

    # DAY_PATTERN = re.compile(r"\bD(\d{1,2})\b", re.IGNORECASE)

    def __init__(self, root_directory, session):
        self.session = session
        self.root_directory = root_directory
        self.TARGET_PREFIX = f"_25 AEP Hometown Week Itinerary Workbook - S{self.session} - "
        self.SESSION_FOLDER = f"Session {self.session} itineraries"

    def find_range(self, sheet, start_row=8):
        for row in range(start_row, sheet.max_row + 1):
            cell_val = sheet.cell(row=row, column=1).value
            val = str(cell_val).strip().lower() if cell_val else ""
            if "today's resource library activities" in val:
                return start_row, row - 1
        return start_row, sheet.max_row

    def extract_data(self, sheet, start_row, end_row, end_col=8):
        data = []
        for row in range(start_row, end_row + 1):
            row_data = []
            for col in range(1, end_col + 1):
                val = sheet.cell(row=row, column=col).value
                if isinstance(val, datetime.time):
                    val = val.strftime("%I:%M %p").lstrip("0")  # e.g., 1:30 PM
                elif isinstance(val, datetime.datetime):
                    val = val.strftime("%I:%M %p").lstrip("0")
                row_data.append(val)
            data.append(row_data)
        return data

    def render_image(self, data, headers, output_path):
        import pandas as pd
        import matplotlib.pyplot as plt
        df = pd.DataFrame(data, columns=headers)

        # Set max width limit
        max_line_len = 40

        max_lines_per_row = []
        wrapped_data = []
        for row in data:
            new_row = []
            row_max = 1
            overflow_rows = []

            for i, cell in enumerate(row):
                val = str(cell) if cell is not None else ""
                if len(val) > max_line_len:
                    words = val.split()
                    wrapped_lines = []
                    current_line = ""
                    for word in words:
                        if len(current_line) + len(word) + 1 <= max_line_len:
                            current_line += " " + word
                        else:
                            wrapped_lines.append(current_line.strip())
                            current_line = word
                    wrapped_lines.append(current_line.strip())

                    new_row.append(wrapped_lines[0])
                    for extra_line in wrapped_lines[1:]:
                        overflow = ["" for _ in range(len(row))]
                        overflow[i] = extra_line
                        overflow_rows.append(overflow)

                    row_max = max(row_max, len(wrapped_lines))
                else:
                    new_row.append(val)

            wrapped_data.append(new_row)
            wrapped_data.extend(overflow_rows)
            max_lines_per_row.append(row_max + len(overflow_rows))

        # Estimate total required height using per-row scaling
        row_heights = [max(1, lines * 0.6) for lines in max_lines_per_row]
        fig_height = max(1, sum(row_heights))

        fig, ax = plt.subplots(figsize=(12, fig_height))
        ax.axis('off')

        def visual_len(s):
            return max((len(line) for line in str(s).split('\n')), default=1)

        num_columns = len(headers)
        col_lens = [len(headers[i]) for i in range(num_columns)]
        for row in wrapped_data:
            for i, cell in enumerate(row):
                col_lens[i] = max(col_lens[i], visual_len(cell))
        total = sum(col_lens)
        col_widths = [0.7 * (l / total) for l in col_lens]

        table = ax.table(cellText=wrapped_data,
                        colLabels=headers,
                        cellLoc='center',
                        loc='center',
                        colWidths=col_widths)

        table.auto_set_font_size(False)
        table.set_fontsize(11)
        table.scale(1.0, 1.2)

        for (row, col), cell in table.get_celld().items():
            cell.set_linewidth(1.1)
            cell.get_text().set_wrap(True)
            cell._text.set_clip_on(True)

            if row == 0:
                cell.set_fontsize(16)
                cell.set_text_props(weight='bold')
            else:
                val = str(cell.get_text().get_text())
                cell.set_text_props(weight='bold')
                lower_val = val.lower().strip()
                if val.endswith("*") or lower_val in {"opening circle", "closing circle", "breakfast", "lunch", "dinner"}:
                    cell.get_text().set_color("red")

        fig.subplots_adjust(left=0, right=1, top=1, bottom=0)
        plt.savefig(output_path, dpi=300, bbox_inches='tight', pad_inches=0.01)
        plt.close()

    def process_workbook(self, filepath, output_root):
        base_name = os.path.basename(filepath)
        parts = re.split(r" - ", base_name.replace(".xlsx", "").replace(".xls", ""))
        if len(parts) >= 3:
            label = parts[2].strip()
        else:
            label = base_name.replace(".xlsx", "").replace(".xls", "")

        wb = load_workbook(filepath, data_only=True)

        sheetnames = wb.sheetnames
        risk_index = None
        for i, name in enumerate(sheetnames):
            if "risk matrix" in name.lower():
                risk_index = i
                break

        if risk_index is None:
            print("Risk Matrix sheet not found.")
            return

        day_indices = list(range(risk_index - 8, risk_index))
        for offset, i in enumerate(day_indices, start=1):
            if i < 0 or i >= len(sheetnames):
                continue
            sheetname = sheetnames[i]
            day_number = offset
            sheet = wb[sheetname]
            start_row, end_row = self.find_range(sheet)
            data = self.extract_data(sheet, start_row, end_row, end_col=3)

            headers = ["Activity", "Time", "Location"]
            day_folder = os.path.join(output_root, f"Day {day_number}")
            os.makedirs(day_folder, exist_ok=True)
            output_path = os.path.join(day_folder, f"{label} - {sheetname.strip()}.png")
            self.render_image(data, headers, output_path)
            print(f"Saved: {output_path}")

    def traverse_exchange_folders(self, root_dir):
        output_root = os.path.join(root_dir, self.SESSION_FOLDER)
        os.makedirs(output_root, exist_ok=True)

        for exchange_folder in os.listdir(root_dir):
            exchange_path = os.path.join(root_dir, exchange_folder)
            if not os.path.isdir(exchange_path):
                continue
            for filename in os.listdir(exchange_path):
                if "workbook" in filename.lower():
                    filepath = os.path.join(exchange_path, filename)
                    self.process_workbook(filepath, output_root)

    def crop_image(self, path):
        from PIL import Image, ImageOps
        img = Image.open(path).convert("RGB")
        gray = ImageOps.grayscale(img)
        inverted = ImageOps.invert(gray)
        bbox = inverted.getbbox()
        if bbox:
            cropped = img.crop(bbox)
            cropped.save(path)

    def crop_all_images_in_directory(self, root_dir):
        for dirpath, _, filenames in os.walk(root_dir):
            for filename in filenames:
                if filename.endswith(".png"):
                    path = os.path.join(dirpath, filename)
                    self.crop_image(path)
                    print(f"Cropped: {path}")

    def run_itinerary_extraction(self):
        if not self.root_directory:
            print("Root directory not provided. Exiting.")
            return
        root_directory = os.path.abspath(self.root_directory)
        session_path = os.path.join(root_directory, self.SESSION_FOLDER)
        if os.path.exists(session_path):
            for root, dirs, files in os.walk(session_path, topdown=False):
                for name in files:
                    os.remove(os.path.join(root, name))
                for name in dirs:
                    os.rmdir(os.path.join(root, name))
        self.traverse_exchange_folders(root_directory)
        self.crop_all_images_in_directory(os.path.join(root_directory, self.SESSION_FOLDER))
        print("Itinerary exctraction finished.")